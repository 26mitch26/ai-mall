"""Build the project-only resume from the reviewed canonical text. Use bundled python-docx."""
import argparse
from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml.ns import qn


def build(output_dir: Path):
    source = Path(__file__).resolve().parents[1] / 'docs/product/resume-project-section.md'
    text = source.read_text(encoding='utf-8')
    output_dir.mkdir(parents=True, exist_ok=True)
    doc = Document()
    section = doc.sections[0]
    # Preserve the existing Chinese project resume's A4 layout.
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(1.8)
    section.left_margin = section.right_margin = Cm(1.8)
    for name in ('Normal', 'Title', 'List Bullet'):
        style = doc.styles[name]
        style.font.name = 'Microsoft YaHei'
        style._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), '微软雅黑')
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.size = Pt(10.5)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.line_spacing = 1.15
        for border in list(style.element.xpath('.//w:pBdr')):
            border.getparent().remove(border)
    doc.styles['Title'].font.size = Pt(18)
    doc.styles['Title'].font.bold = True
    for block in text.strip().split('\n\n'):
        if block.startswith('# '):
            doc.add_paragraph(block[2:], 'Title')
        elif block.startswith('- '):
            for line in block.splitlines():
                lead, body = line[2:].split('：', 1)
                p = doc.add_paragraph(style='List Bullet')
                p.add_run(lead + '：').bold = True
                p.add_run(body)
        else:
            p = doc.add_paragraph(block)
            if block.startswith(('评测口径：', '范围与口径：', '项目范围：')):
                for run in p.runs:
                    run.font.size = Pt(9)
    path = output_dir / 'AI产品经理项目简历.docx'
    doc.save(path)
    (output_dir / 'AI产品经理项目简历_STAR文字版.md').write_text(text, encoding='utf-8')
    print(path)


if __name__ == '__main__':
    parser = argparse.ArgumentParser()
    parser.add_argument('--output-dir', type=Path, required=True)
    args = parser.parse_args()
    build(args.output_dir)
